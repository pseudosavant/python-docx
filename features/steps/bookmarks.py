"""Acceptance steps for public bookmark creation and lookup."""

from io import BytesIO

from behave import given, then, when

from docx import Document


@given("a document with a bookmark on a heading")
def given_bookmark(context):
    context.document = Document()
    heading = context.document.add_heading("Details")
    context.document.bookmarks.add("Details", paragraph=heading)


@when("I save and reopen the bookmark document")
def when_roundtrip(context):
    stream = BytesIO()
    context.document.save(stream)
    context.document = Document(stream)


@then("I can find its heading by bookmark name")
def then_lookup(context):
    assert context.document.bookmarks.get("details").paragraph.text == "Details"


@when("I try to reuse that bookmark name")
def when_duplicate(context):
    try:
        context.document.bookmarks.add("DETAILS", paragraph=context.document.add_paragraph("Other"))
    except ValueError:
        context.rejected = True
    else:
        context.rejected = False


@then("the original bookmark target is preserved")
def then_preserved(context):
    assert context.rejected
    assert len(context.document.bookmarks) == 1
    assert context.document.bookmarks.get("Details").paragraph.text == "Details"
