"""Acceptance steps for DOTX instantiation."""

from io import BytesIO
from pathlib import Path

from behave import given, then, when

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT


@given("a Word-authored DOTX template with styles and content")
def given_dotx(context):
    path = Path(__file__).parents[2] / "tests" / "test_files" / "rich-template-word.dotx"
    context.template_bytes = path.read_bytes()


@when("I instantiate a document from that template and save it")
def when_dotx_saved(context):
    document = Document(BytesIO(context.template_bytes))
    document.add_paragraph("Generated report", "Brand Text")
    stream = BytesIO()
    document.save(stream)
    context.document = Document(stream)


@then("the saved package is a DOCX and retains its template content")
def then_dotx_instantiated(context):
    document = context.document
    assert document.part.content_type == CT.WML_DOCUMENT_MAIN
    assert document.paragraphs[0].text == "Template heading"
    assert document.paragraphs[-1].text == "Generated report"
    assert document.sections[0].header.paragraphs[0].text == "Brand header"
    assert len(document.inline_shapes) == 1
