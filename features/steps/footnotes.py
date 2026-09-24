"""Acceptance steps for native footnote creation."""

from io import BytesIO

from behave import given, then, when

from docx import Document


@given("a document with a native footnote between existing runs")
def given_footnote(context):
    context.document = Document()
    paragraph = context.document.add_paragraph("Statement")
    before = paragraph.runs[0]
    paragraph.add_run(" continues.")
    note = context.document.add_footnote(before, "Explanation")
    note.paragraphs[0].add_run(" with emphasis").italic = True
    note.add_paragraph("More detail")


@when("I save and reopen the footnote document")
def when_footnote_roundtrip(context):
    stream = BytesIO()
    context.document.save(stream)
    context.document = Document(stream)


@then("the reference position and formatted note content are preserved")
def then_footnote_preserved(context):
    document = context.document
    runs = document.paragraphs[0].runs
    assert [run.text for run in runs] == ["Statement", "", " continues."]
    note = document.footnotes.get(runs[1].footnote_ids[0])
    assert note.text == " Explanation with emphasis\nMore detail"
    assert note.paragraphs[0].runs[-1].italic
