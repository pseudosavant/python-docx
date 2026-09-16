"""Acceptance steps for hyperlink authoring."""

from io import BytesIO

from behave import given, then, when
from behave.runner import Context

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.text.paragraph import Paragraph

from helpers import test_docx, test_file


def authoring_paragraph(document: DocumentObject, story: str) -> Paragraph:
    if story == "body":
        return document.paragraphs[0]
    if story == "cell":
        return document.tables[0].cell(0, 0).paragraphs[0]
    if story == "header":
        return document.sections[0].header.paragraphs[0]
    if story == "footer":
        return document.sections[0].footer.paragraphs[0]
    raise ValueError(f"Unknown story: {story}")


@given("a hyperlink authoring paragraph in a {story}")
def given_a_hyperlink_authoring_paragraph(context: Context, story: str):
    context.document = Document()
    context.document.add_paragraph()
    context.document.add_table(1, 1)
    context.story = story
    context.paragraph = authoring_paragraph(context.document, story)
    context.paragraph.add_run("Before ")


@when("I append an external hyperlink between ordinary runs")
def when_I_append_an_external_hyperlink(context: Context):
    context.paragraph.add_hyperlink("the guide", address="https://example.com/guide")
    context.paragraph.add_run(" after")


@then("the new hyperlink and surrounding text survive saving")
def then_the_new_hyperlink_survives_saving(context: Context):
    stream = BytesIO()
    context.document.save(stream)
    paragraph = authoring_paragraph(Document(stream), context.story)
    assert paragraph.text == "Before the guide after"
    assert [run.text for run in paragraph.runs] == ["Before ", " after"]
    assert paragraph.hyperlinks[0].url == "https://example.com/guide"


@when("I append a hyperlink to {address}")
def when_I_append_a_hyperlink_to_an_address(context: Context, address: str):
    context.address = address
    context.paragraph.add_hyperlink("label", address=address)


@then("the supplied hyperlink destination survives saving unchanged")
def then_the_destination_survives_saving(context: Context):
    stream = BytesIO()
    context.document.save(stream)
    hyperlink = Document(stream).paragraphs[0].hyperlinks[0]
    assert hyperlink.address == hyperlink.url == context.address
    assert hyperlink.fragment == ""


@when("I create a hyperlink without an initial label")
def when_I_create_a_hyperlink_without_a_label(context: Context):
    context.hyperlink = context.paragraph.add_hyperlink(address="guide.pdf")
    assert context.hyperlink.runs == []


@then("I can build its label by appending runs")
def then_I_can_build_its_label_by_appending_runs(context: Context):
    context.hyperlink.add_run("the ")
    context.hyperlink.add_run("guide").bold = True
    assert context.hyperlink.text == "the guide"


@when("I try to create a hyperlink with an empty address")
def when_I_try_to_create_a_hyperlink_with_an_empty_address(context: Context):
    context.original_relationships = dict(context.paragraph.part.rels)
    try:
        context.paragraph.add_hyperlink("label", address="")
    except ValueError:
        return
    raise AssertionError("Expected a ValueError for the empty address")


@then("no hyperlink or relationship has been added")
def then_no_hyperlink_or_relationship_has_been_added(context: Context):
    assert context.paragraph.text == "Before "
    assert context.paragraph.hyperlinks == []
    assert dict(context.paragraph.part.rels) == context.original_relationships


@given("an existing hyperlink for authoring")
def given_an_existing_hyperlink_for_authoring(context: Context):
    context.document = Document(test_docx("par-hyperlinks"))
    context.document.styles.add_style("Link emphasis", WD_STYLE_TYPE.CHARACTER)
    context.hyperlink = context.document.paragraphs[1].hyperlinks[0]
    context.original_text = context.hyperlink.text
    context.original_run_count = len(context.hyperlink.runs)


@when("I append formatted runs to the hyperlink")
def when_I_append_formatted_runs_to_the_hyperlink(context: Context):
    context.hyperlink.add_run(" Café\t").bold = True
    context.hyperlink.add_run("code\n", "Link emphasis").font.name = "Consolas"


@then("the appended runs retain their text and formatting after saving")
def then_the_appended_runs_retain_their_text_and_formatting(context: Context):
    stream = BytesIO()
    context.document.save(stream)
    hyperlink = Document(stream).paragraphs[1].hyperlinks[0]
    assert hyperlink.text == context.original_text + " Café\tcode\n"
    assert hyperlink.runs[-2].bold is True
    assert hyperlink.runs[-1].style.name == "Link emphasis"
    assert hyperlink.runs[-1].font.name == "Consolas"


@when("I append a picture run to the hyperlink")
def when_I_append_a_picture_run_to_the_hyperlink(context: Context):
    context.picture = context.hyperlink.add_run().add_picture(test_file("python-icon.jpeg"))


@then("the hyperlink contains the picture after saving")
def then_the_hyperlink_contains_the_picture_after_saving(context: Context):
    stream = BytesIO()
    context.document.save(stream)
    hyperlink = Document(stream).paragraphs[1].hyperlinks[0]
    assert hyperlink.text == context.original_text
    assert len(hyperlink.runs) == context.original_run_count + 1
    assert len(list(hyperlink.runs[-1].iter_inner_content())) == 1
    assert context.picture.width > 0


@when("I try to append a run with a missing character style")
def when_I_try_to_append_a_run_with_a_missing_character_style(context: Context):
    try:
        context.hyperlink.add_run("label", "Missing hyperlink style")
    except KeyError:
        return
    raise AssertionError("Expected a KeyError for the missing style")


@then("the hyperlink's runs remain unchanged")
def then_the_hyperlinks_runs_remain_unchanged(context: Context):
    assert context.hyperlink.text == context.original_text
    assert len(context.hyperlink.runs) == context.original_run_count
