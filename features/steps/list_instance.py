"""Acceptance steps for public numbering operations."""

from io import BytesIO

from behave import given, then, when

from docx import Document
from docx.shared import Inches


@given("a document with an independent list starting at three")
def given_independent_list(context):
    context.document = Document()
    context.sequence = context.document.add_list(start=3)


@when("I add two items with an intervening paragraph")
def when_add_items(context):
    context.sequence.apply(context.document.add_paragraph("Third", "List Number"))
    context.document.add_paragraph("Intervening text")
    context.sequence.apply(context.document.add_paragraph("Fourth", "List Number"))


@when("I restart the list at one")
def when_restart(context):
    restarted = context.sequence.restart()
    restarted.apply(context.document.add_paragraph("New first", "List Number"))


@when("I add an item with an unnumbered continuation")
def when_add_continuation(context):
    context.sequence.apply(context.document.add_paragraph("Third", "List Number"))
    context.sequence.apply_continuation(context.document.add_paragraph("More third", "List Number"))


@when("I save and reopen the list document")
def when_roundtrip(context):
    stream = BytesIO()
    context.document.save(stream)
    context.document = Document(stream)


@then("the items retain their separate numbering identities")
def then_numbering(context):
    paragraphs = context.document.paragraphs
    first = paragraphs[0]._p.pPr.numPr.numId.val
    assert paragraphs[2]._p.pPr.numPr.numId.val == first
    assert paragraphs[3]._p.pPr.numPr.numId.val != first
    numbering = context.document.part.numbering_part.element
    assert numbering.num_having_numId(first).lvlOverride_lst[0].startOverride.val == 3


@then("the continuation retains text alignment without a number")
def then_continuation(context):
    paragraph = context.document.paragraphs[1]
    assert paragraph._p.pPr.numPr.numId.val == 0
    assert paragraph.paragraph_format.left_indent == Inches(0.25)
    assert paragraph.paragraph_format.first_line_indent == 0
