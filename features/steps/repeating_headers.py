"""Acceptance steps for configurable repeating headers."""

from io import BytesIO

from behave import given, then, when

from docx import Document


@given("a table with two leading header rows")
def given_headers(context):
    context.document = Document()
    table = context.document.add_table(rows=3, cols=1)
    for row in table.rows[:2]:
        row.repeat_as_header = True


@when("I save and reopen the repeating-header document")
def when_headers_roundtrip(context):
    stream = BytesIO()
    context.document.save(stream)
    context.document = Document(stream)


@then("both leading rows remain headers and the body remains unmarked")
def then_headers_preserved(context):
    table = context.document.tables[0]
    assert [row.repeat_as_header for row in table.rows] == [True, True, None]
    table.rows[0].repeat_as_header = False
    assert table.rows[0].repeat_as_header is False
    table.rows[0].repeat_as_header = None
    assert table.rows[0].repeat_as_header is None
