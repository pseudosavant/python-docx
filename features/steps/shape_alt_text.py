"""Acceptance steps for inline shape alternative text."""

from io import BytesIO
from pathlib import Path

from behave import given, then, when

from docx import Document


@given("a picture with no alternative text")
def given_picture_without_alternative_text(context):
    context.document = Document()
    picture = Path(__file__).parents[2] / "tests" / "test_files" / "monty-truth.png"
    context.picture = context.document.add_picture(str(picture))
    assert context.picture.description is None
    assert context.picture.title is None


@when("I set its description and title")
def when_set_picture_alternative_text(context):
    context.picture.description = 'Café chart & "résumé"'
    context.picture.title = "Revenue"


@when("I empty its description and remove its title")
def when_clear_picture_alternative_text(context):
    context.picture.description = ""
    context.picture.title = None


@when("I save and reopen the picture document")
def when_reopen_picture_document(context):
    stream = BytesIO()
    context.document.save(stream)
    context.document = Document(stream)
    context.picture = context.document.inline_shapes[0]


@then("its description and title are preserved")
def then_picture_alternative_text_is_preserved(context):
    assert context.picture.description == 'Café chart & "résumé"'
    assert context.picture.title == "Revenue"


@then("its description is empty and its title is absent")
def then_picture_alternative_text_is_cleared(context):
    assert context.picture.description == ""
    assert context.picture.title is None
