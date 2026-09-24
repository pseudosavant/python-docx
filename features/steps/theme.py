"""Acceptance tests for document theme fonts."""

from io import BytesIO

from behave import given, then, when

from docx import Document


@given("a document with theme-based heading and body styles")
def given_theme_based_styles(context):
    document = context.document = Document()
    document.styles.default_font.theme_font = "minor"
    document.styles["Normal"].font.theme_font = "minor"
    document.styles["Heading 1"].font.theme_font = "major"
    document.styles["Heading 1"].linked_style.font.theme_font = "major"
    document.add_heading("Heading")
    document.add_paragraph("Body")


@when("I set its theme fonts to Aptos Display and Aptos and reopen it")
def when_set_theme_and_reopen(context):
    document = context.document
    document.theme_fonts.major_latin = "Aptos Display"
    document.theme_fonts.minor_latin = "Aptos"
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    context.document = Document(stream)


@then("its heading and body styles still reference the theme")
def then_styles_reference_theme(context):
    styles = context.document.styles
    assert styles.default_font.theme_font == "minor"
    assert styles["Normal"].font.theme_font == "minor"
    assert styles["Heading 1"].font.theme_font == "major"
    assert styles["Heading 1"].linked_style.font.theme_font == "major"
    assert styles["Normal"].font.name is None
    assert styles["Heading 1"].font.name is None


@then("its theme fonts are Aptos Display and Aptos")
def then_theme_typefaces(context):
    fonts = context.document.theme_fonts
    assert fonts.major_latin == "Aptos Display"
    assert fonts.minor_latin == "Aptos"
