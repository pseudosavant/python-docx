"""Check fork identity, release tags, and the installed wheel before publishing."""

from __future__ import annotations

import argparse
import ast
import base64
import os
from importlib import metadata
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

PACKAGE_NAME = "ps-python-docx"
ROOT = Path(__file__).resolve().parents[1]


def source_version() -> str:
    module = ast.parse((ROOT / "src" / "docx" / "__init__.py").read_text(encoding="utf-8"))
    for statement in module.body:
        if isinstance(statement, ast.Assign) and any(
            isinstance(target, ast.Name) and target.id == "__version__"
            for target in statement.targets
        ):
            version = ast.literal_eval(statement.value)
            if isinstance(version, str):
                return version
    raise SystemExit("docx.__version__ must be a literal version string.")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--installed", action="store_true")
    args = parser.parse_args()
    version = source_version()
    tag = os.environ.get("RELEASE_TAG")
    if tag and tag != f"v{version}":
        raise SystemExit(f"Release tag {tag!r} must equal 'v{version}'.")

    installed = metadata.distribution(PACKAGE_NAME)
    if installed.metadata["Name"] != PACKAGE_NAME or installed.version != version:
        raise SystemExit("Distribution name or version does not match the fork source.")
    try:
        metadata.distribution("python-docx")
    except metadata.PackageNotFoundError:
        pass
    else:
        raise SystemExit("Upstream python-docx must not be installed alongside the fork.")

    if args.installed:
        import docx

        if ROOT / "src" in Path(docx.__file__).resolve().parents:
            raise SystemExit("Wheel smoke test imported the source tree instead of the wheel.")
        if docx.__version__ != version:
            raise SystemExit("Installed docx version does not match the distribution version.")
        with TemporaryDirectory() as directory:
            output = Path(directory) / "smoke.docx"
            document = docx.Document()
            document.theme_fonts.major_latin = "Aptos Display"
            document.theme_fonts.minor_latin = "Aptos"
            document.styles.default_font.theme_font = "minor"
            document.styles["Normal"].font.theme_font = "minor"
            document.styles["Heading 1"].font.theme_font = "major"
            document.styles["Heading 1"].linked_style.font.theme_font = "major"
            document.add_heading("Fork wheel smoke test", level=1)
            document.add_paragraph("Editable body text")
            hyperlink = document.add_paragraph().add_hyperlink(
                address="https://example.com", tooltip="Details"
            )
            hyperlink.add_run("Example ")
            hyperlink.add_run("bold").bold = True
            image = BytesIO(
                base64.b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
                )
            )
            picture = document.paragraphs[-1].add_run().add_picture(image)
            picture.description = "Pixel description"
            picture.title = "Pixel title"
            sequence = document.add_list(start=3)
            sequence.apply(document.add_paragraph("Third item", "List Number"))
            sequence.apply_continuation(document.add_paragraph("More detail", "List Number"))
            sequence.apply(document.add_paragraph("Fourth item", "List Number"))
            document.bookmarks.add("SmokeHeading", paragraph=document.paragraphs[0])
            document.paragraphs[1].add_hyperlink(anchor="SmokeHeading")
            document.save(output)
            reopened = docx.Document(output)
            assert reopened.bookmarks.get("SmokeHeading").paragraph.text == "Fork wheel smoke test"
            assert reopened.paragraphs[1].hyperlinks[0].fragment == "SmokeHeading"
            assert reopened.theme_fonts.major_latin == "Aptos Display"
            assert reopened.theme_fonts.minor_latin == "Aptos"
            assert reopened.styles.default_font.theme_font == "minor"
            assert reopened.styles["Heading 1"].font.theme_font == "major"
            assert reopened.styles["Heading 1"].linked_style.font.theme_font == "major"
            saved_link = reopened.paragraphs[2].hyperlinks[0]
            assert saved_link.url == "https://example.com"
            assert saved_link.tooltip == "Details"
            assert saved_link.runs[1].bold is True
            assert reopened.inline_shapes[0].description == "Pixel description"
            assert reopened.inline_shapes[0].title == "Pixel title"
            if [paragraph.text for paragraph in reopened.paragraphs] != [
                "Fork wheel smoke test",
                "Editable body text",
                "Example bold",
                "Third item",
                "More detail",
                "Fourth item",
            ]:
                raise SystemExit("Installed wheel did not round-trip document content.")
    print(f"Verified {PACKAGE_NAME} {version}")


if __name__ == "__main__":
    main()
