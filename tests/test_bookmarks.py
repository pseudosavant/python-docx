"""Bookmark destinations and saved-package preservation."""

from io import BytesIO
from pathlib import Path

import pytest

from docx import Document
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement


class DescribeBookmarks:
    def it_reads_word_authored_bookmarks_and_preserves_their_targets(self):
        document = Document(str(Path(__file__).parent / "test_files" / "bookmarks-word.docx"))
        assert {bookmark.name for bookmark in document.bookmarks} >= {
            "Intro",
            "Details_2",
            "Details_3",
        }
        assert document.bookmarks.get("Details_2").paragraph.text == "Details"
        assert document.bookmarks.get("Details_3").paragraph.text == "Details"
        originals = [
            marker.get(qn("w:id")) for marker in document.element.iter(qn("w:bookmarkStart"))
        ]
        added = document.bookmarks.add("Additional", paragraph=document.add_paragraph("More"))
        assert added._element.get(qn("w:id")) not in originals
        stream = BytesIO()
        document.save(stream)
        loaded = Document(stream)
        assert loaded.bookmarks.get("Intro").paragraph.text == "Intro text"
        assert loaded.paragraphs[3].hyperlinks[0].fragment == "Details_3"

    def it_reserves_bookmark_names_in_unparsed_note_parts(self):
        from docx.opc.constants import CONTENT_TYPE as CT
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part

        document = Document()
        blob = (
            b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            b'<w:footnote w:id="1"><w:p><w:bookmarkStart w:id="0" w:name="NoteTarget"/>'
            b'<w:bookmarkEnd w:id="0"/></w:p></w:footnote></w:footnotes>'
        )
        part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, blob, document.part.package)
        document.part.relate_to(part, RT.FOOTNOTES)
        paragraph = document.add_paragraph()
        with pytest.raises(ValueError, match="already exists"):
            document.bookmarks.add("NoteTarget", paragraph=paragraph)
        assert document.bookmarks.add("Body", paragraph=paragraph)._element.get(qn("w:id")) == "1"
        assert part.blob == blob

    @pytest.mark.parametrize("cell", [False, True])
    def it_adds_empty_paragraph_targets_and_roundtrips(self, cell):
        document = Document()
        paragraph = (
            document.add_table(1, 1).cell(0, 0).paragraphs[0] if cell else document.add_paragraph()
        )
        paragraph.style = "Heading 1"
        paragraph.add_run("Bold heading").bold = True
        bookmark = document.bookmarks.add("Target", paragraph=paragraph)
        assert bookmark.name == "Target"
        assert bookmark.paragraph.text == "Bold heading"
        assert document.bookmarks.get("target").paragraph.runs[0].bold is True
        assert document.bookmarks.get("absent") is None
        assert len(document.bookmarks) == 1
        stream = BytesIO()
        document.save(stream)
        loaded = Document(stream)
        assert loaded.bookmarks.get("TARGET").paragraph.text == "Bold heading"
        markers = paragraph._p.xpath("./w:bookmarkStart | ./w:bookmarkEnd")
        assert [marker.get(qn("w:id")) for marker in markers] == ["0", "0"]
        assert list(paragraph._p)[1:3] == markers

    def it_wraps_consecutive_runs_without_changing_formatting(self):
        document = Document()
        paragraph = document.add_paragraph()
        first = paragraph.add_run("first")
        second = paragraph.add_run("second")
        second.italic = True
        paragraph.add_run("outside")
        document.bookmarks.add("Range", runs=[first, second])
        assert paragraph.text == "firstsecondoutside"
        assert paragraph.runs[1].italic is True
        assert [child.tag for child in paragraph._p] == [
            qn("w:bookmarkStart"),
            qn("w:r"),
            qn("w:r"),
            qn("w:bookmarkEnd"),
            qn("w:r"),
        ]
        document.bookmarks.add("Nested", runs=first)
        assert len(document.bookmarks) == 2

    @pytest.mark.parametrize(
        "name", ["", "9start", "has space", "has-dash", "a" * 41, "cafÃ©", "bad\x00name"]
    )
    def it_rejects_unsupported_names_before_mutation(self, name):
        document = Document()
        paragraph = document.add_paragraph("unchanged")
        before = document.element.xml
        with pytest.raises(ValueError, match="bookmark names"):
            document.bookmarks.add(name, paragraph=paragraph)
        assert document.element.xml == before

    def it_rejects_duplicate_names_and_reserves_orphaned_ids(self):
        document = Document()
        paragraph = document.add_paragraph()
        orphan = OxmlElement("w:bookmarkEnd")
        orphan.set(qn("w:id"), "0")
        paragraph._p.append(orphan)
        bookmark = document.bookmarks.add("Target", paragraph=paragraph)
        assert bookmark._element.get(qn("w:id")) == "1"
        before = document.element.xml
        with pytest.raises(ValueError, match="already exists"):
            document.bookmarks.add("TARGET", paragraph=paragraph)
        assert document.element.xml == before

    @pytest.mark.parametrize(
        "invalid",
        [
            "empty",
            "reverse",
            "gap",
            "duplicate",
            "other-paragraph",
            "other-document",
            "header",
            "both",
            "neither",
        ],
    )
    def it_rejects_invalid_ranges_before_mutation(self, invalid):
        document = Document()
        paragraph = document.add_paragraph()
        runs = [paragraph.add_run(str(i)) for i in range(3)]
        kwargs = {
            "empty": {"runs": []},
            "reverse": {"runs": runs[::-1]},
            "gap": {"runs": [runs[0], runs[2]]},
            "duplicate": {"runs": [runs[0], runs[0]]},
            "other-paragraph": {"runs": [runs[0], document.add_paragraph().add_run()]},
            "other-document": {"paragraph": Document().add_paragraph()},
            "header": {"paragraph": document.sections[0].header.paragraphs[0]},
            "both": {"paragraph": paragraph, "runs": runs},
            "neither": {},
        }[invalid]
        before = document.element.xml
        with pytest.raises(ValueError, match="supply|runs|paragraph"):
            document.bookmarks.add("Target", **kwargs)
        assert document.element.xml == before

    def it_preserves_existing_markers_in_other_stories(self):
        document = Document()
        paragraph = document.sections[0].header.paragraphs[0]
        marker = OxmlElement("w:bookmarkStart")
        marker.set(qn("w:id"), "0")
        marker.set(qn("w:name"), "HeaderTarget")
        paragraph._p.append(marker)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "0")
        paragraph._p.append(end)
        before = paragraph._p.xml
        target = document.add_paragraph()
        with pytest.raises(ValueError, match="already exists"):
            document.bookmarks.add("headertarget", paragraph=target)
        bookmark = document.bookmarks.add("Body", paragraph=target)
        assert bookmark._element.get(qn("w:id")) == "1"
        assert paragraph._p.xml == before

    def it_preserves_existing_nonportable_names_on_read(self):
        document = Document()
        paragraph = document.add_paragraph()
        marker = OxmlElement("w:bookmarkStart")
        marker.set(qn("w:id"), "0")
        marker.set(qn("w:name"), "cafÃ©")
        paragraph._p.append(marker)
        assert document.bookmarks.get("cafÃ©").name == "cafÃ©"
