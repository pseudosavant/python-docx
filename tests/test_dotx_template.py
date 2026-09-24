"""Instantiate editable DOCX documents from real Word templates."""

from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.package import Package
from docx.parts.document import DocumentPart

FIXTURES = Path(__file__).parent / "test_files"


def changed_content_type(data, replacement):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            blob = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(CT.WML_TEMPLATE_MAIN.encode(), replacement.encode())
            target.writestr(item, blob)
    output.seek(0)
    return output


@pytest.mark.parametrize("stream_input", [False, True])
def it_instantiates_a_document_and_preserves_template_content(tmp_path, stream_input):
    template = FIXTURES / "rich-template-word.dotx"
    original = template.read_bytes()
    raw = Package.open(BytesIO(original))
    assert isinstance(raw.main_document_part, DocumentPart)
    assert raw.main_document_part.content_type == CT.WML_TEMPLATE_MAIN
    document = Document(BytesIO(original) if stream_input else str(template))
    assert document.part.content_type == CT.WML_DOCUMENT_MAIN
    assert document.paragraphs[0].text == "Template heading"
    assert document.styles["Brand Text"].font.name == "Arial"
    assert document.sections[0].left_margin.pt == 63
    assert document.sections[0].header.paragraphs[0].text == "Brand header"
    assert document.sections[0].footer.paragraphs[0].text == "Brand footer"
    assert len(document.inline_shapes) == 1
    source_parts = {str(part.partname): part.blob for part in raw.parts}
    document.add_paragraph("Generated report", "Brand Text")
    output = tmp_path / "report.docx"
    document.save(output)
    reopened = Document(str(output))
    assert reopened.paragraphs[-1].text == "Generated report"
    assert reopened.paragraphs[-1].style.name == "Brand Text"
    assert reopened.part.content_type == CT.WML_DOCUMENT_MAIN
    for part in reopened.part.package.parts:
        if part is not reopened.part:
            assert part.blob == source_parts[str(part.partname)]
    for original_part in raw.parts:
        saved_part = next(
            p for p in reopened.part.package.parts if p.partname == original_part.partname
        )
        assert [
            (r.rId, r.reltype, r.target_ref, r.is_external) for r in saved_part.rels.values()
        ] == [(r.rId, r.reltype, r.target_ref, r.is_external) for r in original_part.rels.values()]
    assert template.read_bytes() == original
    with ZipFile(output) as package:
        content_types = package.read("[Content_Types].xml")
        assert CT.WML_DOCUMENT_MAIN.encode() in content_types
        assert CT.WML_TEMPLATE_MAIN.encode() not in content_types
        assert b"<w:numPr" in package.read("word/document.xml") or b"ListNumber" in package.read(
            "word/styles.xml"
        )


def it_accepts_blank_word_templates_and_preserves_numbering_and_theme():
    template = FIXTURES / "blank-template-word.dotx"
    raw = Package.open(str(template)).main_document_part
    document = Document(str(template))
    for kind in (RT.NUMBERING, RT.THEME, RT.STYLES):
        assert document.part.part_related_by(kind).blob == raw.part_related_by(kind).blob
    assert not any(paragraph.text for paragraph in document.paragraphs)


def it_detects_template_content_independently_of_input_filename(tmp_path):
    path = tmp_path / "template.data"
    path.write_bytes((FIXTURES / "blank-template-word.dotx").read_bytes())
    document = Document(str(path))
    assert document.part.content_type == CT.WML_DOCUMENT_MAIN
    output = tmp_path / "output.dotx"
    document.save(output)
    assert Package.open(str(output)).main_document_part.content_type == CT.WML_DOCUMENT_MAIN


@pytest.mark.parametrize(
    "content_type",
    [
        "application/vnd.ms-word.document.macroEnabled.main+xml",
        "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
        "application/octet-stream",
    ],
)
def it_rejects_macro_enabled_and_unknown_main_part_types(content_type):
    source = (FIXTURES / "blank-template-word.dotx").read_bytes()
    with pytest.raises(ValueError, match="not a Word file"):
        Document(changed_content_type(source, content_type))


def it_preserves_existing_docx_behavior():
    document = Document()
    document.add_paragraph("Normal document")
    output = BytesIO()
    document.save(output)
    reopened = Document(output)
    assert reopened.part.content_type == CT.WML_DOCUMENT_MAIN
    assert reopened.paragraphs[0].text == "Normal document"
