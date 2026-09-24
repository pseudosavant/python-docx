"""Native note authoring and existing Word content round trips."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement


def roundtrip(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return Document(stream)


def it_reading_empty_collection_does_not_create_part():
    document = Document()
    parts = tuple(document.part.package.parts)
    assert len(document.footnotes) == 0
    assert list(document.footnotes) == []
    assert document.footnotes.get(1) is None
    assert tuple(document.part.package.parts) == parts
    with pytest.raises(KeyError):
        document.part.part_related_by(RT.FOOTNOTES)


def it_reference_is_inserted_between_runs_without_changing_them():
    document = Document()
    paragraph = document.add_paragraph()
    before = paragraph.add_run("before")
    before.bold = True
    after = paragraph.add_run("after")
    after.italic = True
    note = document.add_footnote(before, "Café\nSecond paragraph")
    assert note.footnote_id == 1
    assert before.text == "before"
    assert before.bold
    assert after.text == "after"
    assert after.italic
    reopened = roundtrip(document)
    runs = reopened.paragraphs[0].runs
    assert [run.text for run in runs] == ["before", "", "after"]
    assert [run.footnote_ids for run in runs] == [(), (1,), ()]
    assert runs[1].style.name == "Footnote Reference"
    assert runs[1].style.font.superscript is True
    note = reopened.footnotes.get(1)
    assert note.text == " Café\nSecond paragraph"
    assert all(p.style.name == "Footnote Text" for p in note.paragraphs)
    assert note.paragraphs[0]._p.xpath("./w:r/w:footnoteRef")
    assert len(reopened.footnotes) == 1
    assert reopened.footnotes.get(-1) is None
    part = reopened.part.part_related_by(RT.FOOTNOTES)
    assert [(entry.id, entry.type) for entry in part.element.footnote_lst] == [
        (-1, "separator"),
        (0, "continuationSeparator"),
        (1, None),
    ]
    assert reopened.settings.element.xpath("./w:footnotePr/w:footnote/@w:id") == ["-1", "0"]


def it_empty_note_and_rich_content_in_table_cell():
    document = Document()
    reference = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run()
    note = document.add_footnote(reference)
    assert note.text == " "
    note.paragraphs[0].add_run("bold").bold = True
    note.add_paragraph("Second")
    note.paragraphs[1].add_run(" italic").italic = True
    reopened = roundtrip(document)
    assert reopened.tables[0].cell(0, 0).paragraphs[0].runs[1].footnote_ids == (1,)
    note = reopened.footnotes.get(1)
    assert note.text == " bold\nSecond italic"
    assert note.paragraphs[0].runs[-1].bold
    assert note.paragraphs[1].runs[-1].italic


def it_existing_word_notes_and_formatting_are_preserved():
    document = Document(Path(__file__).parent / "test_files" / "footnotes-word.docx")
    old = document.footnotes.get(1)
    assert old.text == " First note text.\nSecond note paragraph."
    assert any(run.bold for run in old.paragraphs[0].runs)
    old_xml = old._element.xml
    new = document.add_footnote(document.add_paragraph("More").runs[0], "New note")
    assert new.footnote_id == 2
    reopened = roundtrip(document)
    assert reopened.footnotes.get(1)._element.xml == old_xml
    assert reopened.footnotes.get(2).text == " New note"


@pytest.mark.parametrize("location", ["other", "header", "footer", "note", "detached", "nested"])
def it_invalid_locations_leave_package_unchanged(location):
    document = Document()
    if location == "other":
        reference = Document().add_paragraph("Other").runs[0]
    elif location in ("header", "footer"):
        reference = getattr(document.sections[0], location).paragraphs[0].add_run("Text")
    elif location == "note":
        note = document.add_footnote(document.add_paragraph("One").runs[0])
        reference = note.paragraphs[0].add_run("Nested")
    else:
        paragraph = document.add_paragraph("Text")
        reference = paragraph.runs[0]
        if location == "detached":
            paragraph._p.getparent().remove(paragraph._p)
        else:
            hyperlink = OxmlElement("w:hyperlink")
            paragraph._p.append(hyperlink)
            hyperlink.append(reference._r)
    before = {str(part.partname): part.blob for part in document.part.package.parts}
    with pytest.raises(ValueError, match="attached body or table-cell"):
        document.add_footnote(reference, "Invalid")
    assert {str(part.partname): part.blob for part in document.part.package.parts} == before


@pytest.mark.parametrize("text", [None, 5, "invalid\x00text"])
def it_invalid_text_does_not_leave_orphaned_note(text):
    document = Document()
    run = document.add_paragraph("Text").runs[0]
    before = {str(part.partname): part.blob for part in document.part.package.parts}
    with pytest.raises((TypeError, ValueError)):
        document.add_footnote(run, text)
    assert {str(part.partname): part.blob for part in document.part.package.parts} == before


def it_existing_note_styles_are_not_overwritten():
    document = Document()
    styles = document.styles
    text_style = styles.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
    text_style.font.italic = True
    reference_style = styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
    reference_style.font.bold = True
    before = (text_style.element.xml, reference_style.element.xml)
    document.add_footnote(document.add_paragraph("Text").runs[0], "Note")
    assert (text_style.element.xml, reference_style.element.xml) == before


def it_identifiers_avoid_dangling_references_and_each_call_creates_new_note():
    document = Document()
    paragraph = document.add_paragraph("Text")
    marker = OxmlElement("w:footnoteReference")
    marker.set(qn("w:id"), "20")
    paragraph.add_run()._r.append(marker)
    first = document.add_footnote(paragraph.runs[0], "First")
    second = document.add_footnote(paragraph.runs[-1], "Second")
    assert (first.footnote_id, second.footnote_id) == (21, 22)
    assert document.footnotes.get(20) is None
    assert paragraph.runs[-2].footnote_ids == (20,)


def it_package_content_types_and_relationships():
    document = Document()
    document.add_footnote(document.add_paragraph("Text").runs[0], "Note")
    stream = BytesIO()
    document.save(stream)
    with ZipFile(stream) as package:
        assert b"wordprocessingml.footnotes+xml" in package.read("[Content_Types].xml")
        assert b"/footnotes" in package.read("word/_rels/document.xml.rels")


def it_rejects_wrong_argument_and_incompatible_style_before_mutation():
    document = Document()
    with pytest.raises(TypeError, match="Run"):
        document.add_footnote(None)
    document.styles.add_style("Footnote Text", WD_STYLE_TYPE.CHARACTER)
    run = document.add_paragraph("Text").runs[0]
    before = {str(part.partname): part.blob for part in document.part.package.parts}
    with pytest.raises(ValueError, match="incompatible style"):
        document.add_footnote(run, "Note")
    assert {str(part.partname): part.blob for part in document.part.package.parts} == before


def it_uses_an_available_id_when_the_largest_identifier_is_exhausted():
    document = Document()
    run = document.add_paragraph("Text").runs[0]
    note = document.add_footnote(run, "Existing")
    note._element.id = 2**31 - 1
    assert document.add_footnote(run, "New").footnote_id == 2


def it_preserves_note_numbering_settings_when_adding_separators():
    document = Document()
    props = document.settings.element.get_or_add_footnotePr()
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "lowerRoman")
    props.append(number_format)
    document.add_footnote(document.add_paragraph("Text").runs[0], "Note")
    assert document.settings.element.xpath("./w:footnotePr/w:numFmt/@w:val") == ["lowerRoman"]
    assert len(document.settings.element.xpath("./w:footnotePr")) == 1
