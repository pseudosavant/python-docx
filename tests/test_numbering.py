"""Numbering sequences, template preservation, and continuation paragraphs."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Inches


def num_id(paragraph):
    return paragraph._p.pPr.numPr.numId.val


class DescribeListInstance:
    def it_uses_a_word_authored_multilevel_template(self):
        document = Document(str(Path(__file__).parent / "test_files" / "numbering-word.docx"))
        sequence = document.add_list("Example outline", start=3)
        assert sequence.levels == tuple(range(9))
        for text, level in (("Third", 0), ("Child of third", 1), ("Fourth", 0)):
            paragraph = document.add_paragraph(text, "Example outline")
            sequence.apply(paragraph, level=level)
        continuation = document.add_paragraph("More child", "Example outline")
        sequence.apply_continuation(continuation, level=1)
        assert continuation.paragraph_format.left_indent == Inches(0.5)
        assert continuation.paragraph_format.first_line_indent == 0
        stream = BytesIO()
        document.save(stream)
        loaded = Document(stream)
        assert [p._p.pPr.numPr.ilvl.val for p in loaded.paragraphs[3:6]] == [0, 1, 0]

    def it_preserves_independent_sequences_when_the_original_is_resumed(self):
        document = Document()
        original = document.add_list(start=3)
        original.apply(document.add_paragraph("Third", "List Number"))
        original.apply(document.add_paragraph("Fourth", "List Number"))
        restarted = original.restart()
        restarted.apply(document.add_paragraph("New first", "List Number"))
        original.apply(document.add_paragraph("Fifth", "List Number"))
        # Separate abstract definitions prevent Word from displaying 2 for Fifth.
        assert original._num.abstractNumId.val != restarted._num.abstractNumId.val
        assert original._abstract.find(qn("w:nsid")).get(qn("w:val")) != restarted._abstract.find(
            qn("w:nsid")
        ).get(qn("w:val"))
        assert num_id(document.paragraphs[0]) == num_id(document.paragraphs[3])

    def it_starts_restarts_and_resumes_without_mutating_shared_definitions(self):
        document = Document()
        numbering = document.part.numbering_part.element
        abstracts = [deepcopy(a) for a in numbering.findall(qn("w:abstractNum"))]
        sequence = document.add_list(start=3)
        first = document.add_paragraph("Third", "List Number")
        sequence.apply(first)
        document.add_paragraph("Intervening text")
        second = document.add_paragraph("Fourth", "List Number")
        sequence.apply(second)
        restarted = sequence.restart(start=8)
        third = document.add_paragraph("Eighth", "List Number")
        restarted.apply(third)
        assert num_id(first) == num_id(second)
        assert num_id(first) != num_id(third)
        assert sequence._num.lvlOverride_lst[0].startOverride.val == 3
        assert restarted._num.lvlOverride_lst[0].startOverride.val == 8
        from lxml import etree

        assert [etree.tostring(a) for a in abstracts] == [
            etree.tostring(deepcopy(a))
            for a in numbering.findall(qn("w:abstractNum"))[: len(abstracts)]
        ]
        stream = BytesIO()
        document.save(stream)
        loaded = Document(stream)
        assert [
            num_id(p) for p in [loaded.paragraphs[0], loaded.paragraphs[2], loaded.paragraphs[3]]
        ] == [num_id(first), num_id(second), num_id(third)]

    def it_preserves_style_inheritance_and_numbers_each_list_independently(self):
        document = Document()
        style = document.styles.add_style("Custom numbered", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["List Number 2"]
        first = document.add_list(style, start=0)
        second = document.add_list(style)
        assert first.default_level == 0
        assert first.levels == (0,)
        assert first._num.abstractNumId.val != second._num.abstractNumId.val
        assert first._num.numId != second._num.numId

    def it_makes_continuations_without_consuming_numbers_and_preserves_alignment(self):
        document = Document()
        sequence = document.add_list("List Number 2")
        item = document.add_paragraph("First", "List Number 2")
        sequence.apply(item)
        continuation = document.add_paragraph("More", "List Number 2")
        sequence.apply_continuation(continuation)
        assert num_id(continuation) == 0
        assert continuation.paragraph_format.left_indent == Inches(0.5)
        assert continuation.paragraph_format.first_line_indent == 0
        assert continuation.paragraph_format.tab_stops[0].position == Inches(0.5)
        assert continuation.style == item.style
        continuation.paragraph_format.left_indent = Inches(1)
        continuation.paragraph_format.first_line_indent = Inches(-0.25)
        sequence.apply_continuation(continuation)
        assert continuation.paragraph_format.left_indent == Inches(1)
        assert continuation.paragraph_format.first_line_indent == 0

    def it_preserves_multilevel_definitions_and_existing_overrides(self):
        document = Document()
        sequence = document.add_list()
        level = deepcopy(sequence._abstract.find(qn("w:lvl")))
        level.set(qn("w:ilvl"), "1")
        level.find(qn("w:lvlText")).set(qn("w:val"), "%1.%2.")
        sequence._abstract.append(level)
        override = sequence._num.add_lvlOverride(1)
        override.add_startOverride(4)
        custom = deepcopy(level)
        custom.find(qn("w:pPr")).find(qn("w:ind")).set(qn("w:left"), "1440")
        override.append(custom)
        restarted = sequence.restart(start=7)
        assert restarted.levels == (0, 1)
        assert restarted._num.lvlOverride_lst[1].startOverride.val == 4
        outer = document.add_paragraph("Outer")
        inner = document.add_paragraph("Inner")
        after = document.add_paragraph("After")
        restarted.apply(outer)
        restarted.apply(inner, level=1)
        restarted.apply(after)
        assert [p._p.pPr.numPr.ilvl.val for p in (outer, inner, after)] == [0, 1, 0]
        continuation = document.add_paragraph("More inner")
        restarted.apply_continuation(continuation, level=1)
        assert continuation.paragraph_format.left_indent == Inches(1)

    def it_supports_cell_paragraphs(self):
        document = Document()
        paragraph = document.add_table(1, 1).cell(0, 0).paragraphs[0]
        document.add_list("List Bullet").apply(paragraph)
        assert num_id(paragraph) > 0

    @pytest.mark.parametrize("level", [-1, 9, 1, True, 1.0, "1"])
    def it_rejects_invalid_or_undefined_levels_without_mutation(self, level):
        document = Document()
        sequence = document.add_list()
        paragraph = document.add_paragraph("Unchanged")
        before = document.part.numbering_part.element.xml
        paragraph_before = paragraph._p.xml
        for action in (sequence.apply, sequence.apply_continuation):
            with pytest.raises((ValueError, TypeError)):
                action(paragraph, level=level)
            assert paragraph._p.xml == paragraph_before
        with pytest.raises((ValueError, TypeError)):
            sequence.restart(level=level)
        with pytest.raises((ValueError, TypeError)):
            document.add_list(level=level)
        assert document.part.numbering_part.element.xml == before

    @pytest.mark.parametrize("start", [-1, 2147483648, True, "3", 1.5, None])
    def it_validates_start_before_mutation(self, start):
        document = Document()
        before = document.part.numbering_part.element.xml
        with pytest.raises((ValueError, TypeError)):
            document.add_list(start=start)
        assert document.part.numbering_part.element.xml == before

    def it_rejects_foreign_styles_paragraphs_and_other_stories(self):
        document = Document()
        other = Document()
        with pytest.raises(ValueError, match="belong"):
            document.add_list(other.styles["List Number"])
        sequence = document.add_list()
        for paragraph in (other.add_paragraph(), document.sections[0].header.paragraphs[0]):
            before = paragraph._p.xml
            with pytest.raises(ValueError, match="main story"):
                sequence.apply(paragraph)
            with pytest.raises(ValueError, match="main story"):
                sequence.apply_continuation(paragraph)
            assert paragraph._p.xml == before

    def it_rejects_styles_without_numbering_and_cycles(self):
        document = Document()
        with pytest.raises(ValueError, match="numbering definition"):
            document.add_list("Normal")
        with pytest.raises(ValueError, match="paragraph style"):
            document.add_list("Emphasis")
        style = document.styles["List Number"]
        style.base_style = style
        with pytest.raises(ValueError, match="cycle"):
            document.add_list(style)

    @pytest.mark.parametrize(
        "broken", ["missing-part", "missing-num", "missing-abstract", "style-link"]
    )
    def it_rejects_unsupported_template_references_without_mutation(self, broken):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml.parser import OxmlElement

        document = Document()
        numbering = document.part.numbering_part.element
        if broken == "missing-part":
            document.part.drop_rel(
                next(rel.rId for rel in document.part.rels.values() if rel.reltype == RT.NUMBERING)
            )
        elif broken == "missing-num":
            numbering.remove(numbering.num_having_numId(5))
        elif broken == "missing-abstract":
            numbering.remove(numbering.xpath('./w:abstractNum[@w:abstractNumId="7"]')[0])
        else:
            link = OxmlElement("w:numStyleLink")
            link.set(qn("w:val"), "Linked")
            numbering.xpath('./w:abstractNum[@w:abstractNumId="7"]')[0].append(link)
        before = numbering.xml
        with pytest.raises(ValueError, match="numbering"):
            document.add_list()
        assert numbering.xml == before
