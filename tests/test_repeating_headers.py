"""Repeating header settings and preservation of other row properties."""

from io import BytesIO
from pathlib import Path

import pytest

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml.parser import parse_xml
from docx.shared import Inches
from docx.table import _Row


@pytest.mark.parametrize(
    ("setting", "expected"),
    [
        ("", None),
        ("<w:tblHeader/>", True),
        ('<w:tblHeader w:val="1"/>', True),
        ('<w:tblHeader w:val="true"/>', True),
        ('<w:tblHeader w:val="0"/>', False),
        ('<w:tblHeader w:val="false"/>', False),
    ],
)
def it_reads_present_absent_and_explicit_false_states(setting, expected):
    xml = '<w:tr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    row = _Row(parse_xml(f"{xml}<w:trPr>{setting}</w:trPr></w:tr>"), None)
    assert row.repeat_as_header is expected


def it_roundtrips_multiple_leading_rows_without_changing_unrelated_content():
    document = Document()
    table = document.add_table(rows=4, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.rows[0].height = Inches(0.4)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cells_before = [cell._tc.xml for cell in table.rows[0].cells]
    for row in table.rows[:2]:
        row.repeat_as_header = True
        row.repeat_as_header = True
    table.rows[2].repeat_as_header = False
    assert [cell._tc.xml for cell in table.rows[0].cells] == cells_before
    assert len(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")) == 1
    assert [child.tag for child in table.rows[0]._tr.trPr] == [qn("w:trHeight"), qn("w:tblHeader")]
    stream = BytesIO()
    document.save(stream)
    reopened = Document(stream).tables[0]
    assert [row.repeat_as_header for row in reopened.rows] == [True, True, False, None]
    assert reopened.rows[0].height == Inches(0.4)
    assert reopened.rows[0].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
    assert reopened.cell(0, 0).text == "Merged heading"
    reopened.rows[0].repeat_as_header = None
    assert reopened.rows[0].repeat_as_header is None
    assert reopened.rows[0].height == Inches(0.4)
    assert not reopened._tbl.xpath(".//w:cantSplit")


def it_leaves_new_rows_unmarked_and_clearing_an_absent_setting_does_not_add_properties():
    document = Document()
    row = document.add_table(rows=1, cols=1).rows[0]
    before = row._tr.xml
    assert row.repeat_as_header is None
    row.repeat_as_header = None
    assert row._tr.xml == before


@pytest.mark.parametrize("value", ["true", 1, 0, []])
def it_rejects_non_boolean_settings_without_mutation(value):
    row = Document().add_table(rows=1, cols=1).rows[0]
    before = row._tr.xml
    with pytest.raises(TypeError):
        row.repeat_as_header = value
    assert row._tr.xml == before


def it_preserves_a_non_leading_header_flag_without_promising_repetition():
    document = Document()
    table = document.add_table(rows=3, cols=1)
    table.rows[2].repeat_as_header = True
    stream = BytesIO()
    document.save(stream)
    assert [row.repeat_as_header for row in Document(stream).tables[0].rows] == [None, None, True]


def it_reads_word_authored_header_rows():
    document = Document(Path(__file__).parent / "test_files" / "repeating-headers-word.docx")
    assert [row.repeat_as_header for row in document.tables[0].rows[:2]] == [True, True]
    assert document.tables[0].rows[3].repeat_as_header is None
