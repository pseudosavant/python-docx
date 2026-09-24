# pyright: reportPrivateUsage=false

"""Inline shape alternative text and saved-package preservation tests."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import cast
from zipfile import ZipFile

import pytest

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.parser import parse_xml
from docx.oxml.shape import CT_Inline, CT_NonVisualDrawingProps
from docx.shape import InlineShape
from docx.shared import Inches
from docx.text.paragraph import Paragraph

from .unitutil.cxml import element

PICTURE = str(Path(__file__).parent / "test_files" / "monty-truth.png")


def saved_parts(document: DocumentObject) -> dict[str, bytes]:
    stream = BytesIO()
    document.save(stream)
    with ZipFile(stream) as package:
        return {name: package.read(name) for name in package.namelist()}


def paragraph_in(document: DocumentObject, story: str) -> Paragraph:
    if story == "body":
        return document.paragraphs[0]
    if story == "cell":
        return document.tables[0].cell(0, 0).paragraphs[0]
    return getattr(document.sections[0], story).paragraphs[0]


class DescribeInlineShapeAlternativeText:
    @pytest.mark.parametrize(
        ("property_name", "attribute"), [("description", "descr"), ("title", "title")]
    )
    @pytest.mark.parametrize("value", [None, "", 'Café <chart> & "résumé"\n東京\t ', " "])
    def it_reads_sets_and_clears_each_attribute(
        self, property_name: str, attribute: str, value: str | None
    ):
        inline = cast(CT_Inline, element("wp:inline/wp:docPr{id=1,name=Picture 1}"))
        shape = InlineShape(inline)
        assert getattr(shape, property_name) is None
        setattr(shape, property_name, value)
        assert inline.docPr.get(attribute) == value
        reopened = InlineShape(cast(CT_Inline, parse_xml(inline.xml)))
        assert getattr(reopened, property_name) == value
        setattr(shape, property_name, None)
        assert attribute not in inline.docPr.attrib

    @pytest.mark.parametrize("property_name", ["description", "title"])
    @pytest.mark.parametrize("value", [42, False, b"bytes", "invalid\x00text"])
    def it_rejects_invalid_values_without_changing_existing_metadata(
        self, property_name: str, value: object
    ):
        inline = cast(
            CT_Inline, element("wp:inline/wp:docPr{id=1,name=Picture 1,descr=Original,title=Title}")
        )
        shape = InlineShape(inline)
        before = inline.xml
        with pytest.raises((TypeError, ValueError)):
            setattr(shape, property_name, value)
        assert inline.xml == before

    @pytest.mark.parametrize("story", ["body", "cell", "header", "footer"])
    def it_round_trips_independent_descriptions_for_shared_image_bytes(self, story: str):
        document = Document()
        document.add_paragraph("Before ")
        document.add_table(1, 1)
        paragraph = paragraph_in(document, story)
        first = paragraph.add_run().add_picture(PICTURE, width=Inches(1))
        second = paragraph.add_run().add_picture(PICTURE, width=Inches(2))
        paragraph.add_run(" after")
        first.description, first.title = "First description", "First title"
        second.description, second.title = "Second description", "Second title"
        stream = BytesIO()
        document.save(stream)
        reopened = Document(stream)
        drawings = paragraph_in(reopened, story)._p.xpath(".//wp:inline")
        shapes = [InlineShape(inline) for inline in drawings]
        assert [(shape.description, shape.title) for shape in shapes] == [
            ("First description", "First title"),
            ("Second description", "Second title"),
        ]
        assert [shape.width for shape in shapes] == [Inches(1), Inches(2)]
        parts = saved_parts(reopened)
        assert len([name for name in parts if name.startswith("word/media/")]) == 1

    def it_changes_only_the_requested_drawing_attributes(self):
        document = Document()
        shape = document.add_picture(PICTURE, width=Inches(1))
        shape.description, shape.title = "Original", "Original title"
        before = saved_parts(document)
        shape.description = "Updated description"
        assert shape.title == "Original title"
        shape.title = "Updated title"
        assert shape.description == "Updated description"
        after = saved_parts(document)
        assert before.keys() == after.keys()
        assert all(before[name] == after[name] for name in before if name != "word/document.xml")
        expected = parse_xml(before["word/document.xml"])
        props = expected.xpath(".//wp:docPr")[0]
        props.set("descr", "Updated description")
        props.set("title", "Updated title")
        assert parse_xml(after["word/document.xml"]).xml == expected.xml

    def it_preserves_picture_level_metadata_when_updating_the_placed_shape(self):
        document = Document()
        shape = document.add_picture(PICTURE)
        picture_props = cast(
            CT_NonVisualDrawingProps, shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr
        )
        picture_props.set("descr", "Picture metadata")
        picture_props.set("title", "Picture title")
        assert shape.description is None
        assert shape.title is None
        shape.description, shape.title = "Placed shape description", "Placed shape title"
        assert picture_props.get("descr") == "Picture metadata"
        assert picture_props.get("title") == "Picture title"
