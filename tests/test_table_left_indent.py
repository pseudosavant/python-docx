"""Round-trip behavior for direct table indentation."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches


def test_table_left_indent_round_trips_without_changing_other_properties():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.rows[0].cells[0].text = "Header"
    table.rows[1].cells[0].text = "Body"
    table.left_indent = Inches(0.75)

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    reopened_document = Document(stream)
    reopened = reopened_document.tables[0]
    assert reopened.left_indent == Inches(0.75)
    assert reopened.alignment == WD_TABLE_ALIGNMENT.LEFT
    assert reopened.rows[1].cells[0].text == "Body"

    reopened.left_indent = None
    stream = BytesIO()
    reopened_document.save(stream)
    stream.seek(0)
    assert Document(stream).tables[0].left_indent is None
