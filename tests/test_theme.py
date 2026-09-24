"""Theme font behavior verified against saved packages."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn


def saved_parts(document):
    stream = BytesIO()
    document.save(stream)
    with ZipFile(stream) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def reopened(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return Document(stream)


class DescribeThemeFonts:
    def it_reads_and_writes_both_latin_fonts_and_scheme_name(self):
        document = Document()
        before = saved_parts(document)
        fonts = document.theme_fonts
        assert fonts.major_latin == "Calibri"
        assert fonts.minor_latin == "Cambria"
        fonts.name = "Custom fonts"
        fonts.major_latin = "Aptos Display"
        fonts.minor_latin = "Aptos"
        after = saved_parts(document)
        assert after["word/styles.xml"] == before["word/styles.xml"]
        assert after["word/document.xml"] == before["word/document.xml"]
        again = reopened(document).theme_fonts
        assert (again.name, again.major_latin, again.minor_latin) == (
            "Custom fonts",
            "Aptos Display",
            "Aptos",
        )

    @pytest.mark.parametrize("role", ["major", "minor"])
    def it_preserves_other_fonts_colors_and_effects(self, role):
        document = Document()
        before = ET.fromstring(saved_parts(document)["word/theme/theme1.xml"])
        setattr(document.theme_fonts, f"{role}_latin", "Aptos")
        after = ET.fromstring(saved_parts(document)["word/theme/theme1.xml"])
        path = f"{qn('a:themeElements')}/{qn('a:fontScheme')}/{qn(f'a:{role}Font')}/{qn('a:latin')}"
        latin = before.find(path)
        latin.set("typeface", "Aptos")
        for attribute in ("panose", "pitchFamily", "charset"):
            latin.attrib.pop(attribute, None)
        assert ET.canonicalize(ET.tostring(before), strip_text=True) == ET.canonicalize(
            ET.tostring(after), strip_text=True
        )

    def it_keeps_multiple_proxies_synchronized(self):
        document = Document()
        first, second = document.theme_fonts, document.theme_fonts
        first.major_latin = "Georgia"
        assert second.major_latin == "Georgia"
        second.minor_latin = "Verdana"
        assert first.minor_latin == "Verdana"

    def it_creates_a_complete_theme_and_relationship_when_absent(self):
        document = Document()
        rid = next(r.rId for r in document.part.rels.values() if r.reltype == RT.THEME)
        document.part.drop_rel(rid)
        assert "word/theme/theme1.xml" not in saved_parts(document)
        document.theme_fonts.major_latin = "Aptos Display"
        document.theme_fonts.minor_latin = "Aptos"
        again = reopened(document)
        assert again.theme_fonts.major_latin == "Aptos Display"
        assert again.theme_fonts.minor_latin == "Aptos"
        assert sum(r.reltype == RT.THEME for r in again.part.rels.values()) == 1
        root = ET.fromstring(again.part.part_related_by(RT.THEME).blob)
        for tag in ("clrScheme", "fontScheme", "fmtScheme"):
            assert root.find(f"{qn('a:themeElements')}/{qn(f'a:{tag}')}") is not None

    @pytest.mark.parametrize(
        "blob",
        [
            b"<broken",
            b'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>',
        ],
    )
    def it_rejects_invalid_themes_without_replacing_them(self, blob):
        document = Document()
        part = document.part.part_related_by(RT.THEME)
        part._blob = blob
        with pytest.raises(ValueError, match="theme"):
            document.theme_fonts
        assert part._blob == blob

    @pytest.mark.parametrize("role", ["major", "minor"])
    def it_rejects_an_incomplete_font_scheme(self, role):
        document = Document()
        part = document.part.part_related_by(RT.THEME)
        root = ET.fromstring(part.blob)
        scheme = root.find(f"{qn('a:themeElements')}/{qn('a:fontScheme')}")
        font = scheme.find(qn(f"a:{role}Font"))
        font.remove(font.find(qn("a:latin")))
        part._blob = ET.tostring(root)
        with pytest.raises(ValueError, match=f"{role} Latin"):
            document.theme_fonts

    def it_preserves_an_untouched_theme_verbatim(self):
        document = Document()
        part = document.part.part_related_by(RT.THEME)
        original = part.blob
        document.add_paragraph("Body")
        assert saved_parts(document)["word/theme/theme1.xml"] == original

    def it_rejects_a_non_string_typeface_without_mutation(self):
        document = Document()
        fonts = document.theme_fonts
        original = saved_parts(document)["word/theme/theme1.xml"]
        with pytest.raises(TypeError, match="string"):
            fonts.major_latin = None
        assert saved_parts(document)["word/theme/theme1.xml"] == original
