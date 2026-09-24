"""Theme references and style inheritance, independent of Markdown clients."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.styles.styles import Styles
from docx.text.font import Font


class DescribeThemeFontReferences:
    @pytest.mark.parametrize("role", ["major", "minor"])
    def it_replaces_literal_latin_fonts_and_preserves_other_formatting(self, role):
        run = parse_xml(
            '<w:r %s><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" '
            'w:eastAsia="Japanese" w:cs="Arabic" w:eastAsiaTheme="majorEastAsia" '
            'w:cstheme="majorBidi"/><w:b/></w:rPr></w:r>' % nsdecls("w")
        )
        font = Font(run)
        font.theme_font = role
        assert font.name is None
        assert font.theme_font == role
        assert font.bold is True
        fonts = run.rPr.rFonts
        for attr in ("asciiTheme", "hAnsiTheme"):
            assert fonts.get(qn(f"w:{attr}")) == f"{role}HAnsi"
        assert fonts.get(qn("w:eastAsia")) == "Japanese"
        assert fonts.get(qn("w:cs")) == "Arabic"
        assert fonts.get(qn("w:eastAsiaTheme")) == "majorEastAsia"
        assert fonts.get(qn("w:cstheme")) == "majorBidi"

    def it_clears_references_without_removing_literal_names(self):
        run = Document().add_paragraph().add_run("Code")
        run.font.theme_font = "major"
        run.font.name = "Consolas"
        # Existing name assignment semantics remain backward compatible.
        assert run.font.theme_font == "major"
        run.font.theme_font = None
        assert run.font.theme_font is None
        assert run.font.name == "Consolas"

    def it_leaves_unformatted_runs_unchanged_when_cleared(self):
        run = Document().add_paragraph().add_run("Body")
        before = run._r.xml
        assert run.font.theme_font is None
        run.font.theme_font = None
        assert run._r.xml == before

    @pytest.mark.parametrize("value", ["headings", "majorHAnsi", "", False, 1])
    def it_rejects_invalid_values_without_mutation(self, value):
        run = Document().add_paragraph().add_run("Body")
        before = run._r.xml
        with pytest.raises(ValueError, match="theme_font"):
            run.font.theme_font = value
        assert run._r.xml == before

    @pytest.mark.parametrize(
        ("ascii_ref", "hansi_ref", "expected"),
        [
            ("majorAscii", "majorHAnsi", "major"),
            ("minorAscii", "minorHAnsi", "minor"),
            ("majorHAnsi", "minorHAnsi", None),
            ("majorHAnsi", None, None),
            (None, None, None),
        ],
    )
    def it_reads_only_uniform_local_references(self, ascii_ref, hansi_ref, expected):
        run = Document().add_paragraph().add_run("Body")
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.asciiTheme = ascii_ref
        fonts.hAnsiTheme = hansi_ref
        assert run.font.theme_font == expected


class DescribeDefaultAndLinkedFonts:
    @pytest.mark.parametrize(
        "xml",
        [
            "<w:styles %s/>",
            "<w:styles %s><w:docDefaults><w:pPrDefault/></w:docDefaults></w:styles>",
            "<w:styles %s><w:docDefaults><w:rPrDefault/></w:docDefaults></w:styles>",
        ],
    )
    def it_creates_missing_defaults_in_schema_order(self, xml):
        styles = Styles(parse_xml(xml % nsdecls("w")))
        styles.default_font.theme_font = "minor"
        styles.default_font.size = Pt(12)
        assert styles.default_font.theme_font == "minor"
        assert styles.default_font.size == Pt(12)
        defaults = styles.element.find(qn("w:docDefaults"))
        assert defaults[0].tag == qn("w:rPrDefault")
        assert len(defaults.findall(qn("w:rPrDefault"))) == 1

    def it_reads_links_in_both_directions_and_tolerates_missing_targets(self):
        document = Document()
        heading = document.styles["Heading 1"]
        linked = heading.linked_style
        assert linked.name == "Heading 1 Char"
        assert linked.linked_style == heading
        assert document.styles["Normal"].linked_style is None
        heading._element.link.val = "MissingStyle"
        assert heading.linked_style is None

    def it_retains_default_paragraph_properties(self):
        document = Document()
        defaults = document.styles.element.find(qn("w:docDefaults"))
        before = ET.tostring(ET.fromstring(defaults.xml).find(qn("w:pPrDefault")))
        document.styles.default_font.theme_font = "minor"
        after = ET.tostring(ET.fromstring(defaults.xml).find(qn("w:pPrDefault")))
        assert before == after

    def it_round_trips_inheritance_without_direct_run_formatting(self):
        document = Document()
        document.theme_fonts.major_latin = "Aptos Display"
        document.theme_fonts.minor_latin = "Aptos"
        document.styles.default_font.theme_font = "minor"
        document.styles["Normal"].font.theme_font = "minor"
        for level in range(1, 7):
            style = document.styles[f"Heading {level}"]
            style.font.theme_font = "major"
            style.linked_style.font.theme_font = "major"
            document.add_heading(f"Heading {level}", level)
        document.add_paragraph("Body")
        stream = BytesIO()
        document.save(stream)
        stream.seek(0)
        reopened = Document(stream)
        reopened.add_heading("Future heading", 1)
        reopened.add_paragraph("Future body")
        before = BytesIO()
        reopened.save(before)
        reopened.theme_fonts.major_latin = "Georgia"
        reopened.theme_fonts.minor_latin = "Verdana"
        after = BytesIO()
        reopened.save(after)
        with ZipFile(before) as original, ZipFile(after) as changed:
            for name in ("word/styles.xml", "word/document.xml"):
                assert original.read(name) == changed.read(name)
            assert original.read("word/theme/theme1.xml") != changed.read("word/theme/theme1.xml")
            content = ET.fromstring(changed.read("word/document.xml"))
            assert not content.findall(f".//{qn('w:r')}/{qn('w:rPr')}/{qn('w:rFonts')}")
            styles = ET.fromstring(changed.read("word/styles.xml"))
            for level in range(1, 7):
                for style_id in (f"Heading{level}", f"Heading{level}Char"):
                    style = next(s for s in styles if s.get(qn("w:styleId")) == style_id)
                    fonts = style.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
                    assert fonts.get(qn("w:asciiTheme")) == "majorHAnsi"
                    assert fonts.get(qn("w:hAnsiTheme")) == "majorHAnsi"
                    assert fonts.get(qn("w:ascii")) is None
