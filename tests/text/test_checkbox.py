# pyright: reportPrivateUsage=false

"""Saved-document tests for clickable paragraph check boxes."""

from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest

from docx import Document
from docx.oxml.ns import qn
from docx.text.checkbox import CheckBox


@pytest.mark.parametrize(("checked", "glyph"), [(False, "\u2610"), (True, "\u2612")])
def it_round_trips_a_checkbox(checked: bool, glyph: str) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    checkbox = paragraph.add_checkbox(checked)
    paragraph.add_run(" task")
    assert isinstance(checkbox, CheckBox)
    assert checkbox.checked is checked

    stream = BytesIO()
    document.save(stream)
    with ZipFile(stream) as package:
        xml = package.read("word/document.xml")
    assert b"w14:checkbox" in xml
    stream.seek(0)
    reopened = Document(stream)
    paragraph = reopened.paragraphs[0]
    checkbox = paragraph.checkboxes[0]
    assert checkbox.checked is checked
    assert paragraph._p.xpath("./w:sdt/w:sdtContent/w:r/w:t")[0].text == glyph

    checkbox.checked = not checked
    assert checkbox.checked is not checked
    assert paragraph._p.xpath("./w:sdt/w:sdtContent/w:r/w:t")[0].text != glyph
    assert paragraph._p.xpath("./w:sdt/w:sdtPr/w14:checkbox/w14:checked")[0].get(qn("w14:val")) == (
        "0" if checked else "1"
    )


def it_accepts_only_boolean_state() -> None:
    paragraph = Document().add_paragraph()
    with pytest.raises(TypeError, match="checked must be a bool"):
        paragraph.add_checkbox(1)  # type: ignore[arg-type]
    assert paragraph.checkboxes == []
    checkbox = paragraph.add_checkbox()
    with pytest.raises(TypeError, match="checked must be a bool"):
        checkbox.checked = "yes"  # type: ignore[assignment]
    assert checkbox.checked is False


def it_works_in_table_cells() -> None:
    document = Document()
    paragraph = document.add_table(1, 1).cell(0, 0).paragraphs[0]
    paragraph.add_checkbox(True)
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    assert Document(stream).tables[0].cell(0, 0).paragraphs[0].checkboxes[0].checked is True
