"""Test suite for the docx.text.hyperlink module."""

from __future__ import annotations

from typing import cast

import pytest

from docx import types as t
from docx.opc.rel import _Relationship  # pyright: ignore[reportPrivateUsage]
from docx.oxml.text.hyperlink import CT_Hyperlink
from docx.parts.story import StoryPart
from docx.text.hyperlink import Hyperlink
from docx.text.run import Run

from ..unitutil.cxml import element, xml
from ..unitutil.mock import FixtureRequest, Mock, instance_mock, property_mock


class DescribeHyperlink:
    """Unit-test suite for the docx.text.hyperlink.Hyperlink object."""

    @pytest.mark.parametrize(
        ("text", "expected_cxml"),
        [
            (None, "w:hyperlink/w:r"),
            ("", "w:hyperlink/w:r"),
            ("label", 'w:hyperlink/w:r/w:t"label"'),
            (
                " a\tb\nc\rd ",
                'w:hyperlink/w:r/(w:t{xml:space=preserve}" a",w:tab,w:t"b",'
                'w:br,w:t"c",w:br,w:t{xml:space=preserve}"d ")',
            ),
        ],
    )
    def it_can_append_a_run(
        self, text: str | None, expected_cxml: str, fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element("w:hyperlink"))
        hyperlink = Hyperlink(hlink, fake_parent)

        run = hyperlink.add_run(text)

        assert isinstance(run, Run)
        assert run.part is fake_parent.part
        assert hlink.xml == xml(expected_cxml)

    def it_can_apply_a_character_style_to_a_new_run(
        self, request: FixtureRequest, fake_parent: t.ProvidesStoryPart
    ):
        style_prop = property_mock(request, Run, "style")
        hyperlink = Hyperlink(cast(CT_Hyperlink, element("w:hyperlink")), fake_parent)

        hyperlink.add_run("label", "Emphasis")

        style_prop.assert_called_once_with("Emphasis")

    def it_appends_a_run_after_existing_content(self, fake_parent: t.ProvidesStoryPart):
        hlink = cast(CT_Hyperlink, element('w:hyperlink/w:r/w:t"before"'))
        hyperlink = Hyperlink(hlink, fake_parent)

        run = hyperlink.add_run("after")

        assert hlink.xml == xml('w:hyperlink/(w:r/w:t"before",w:r/w:t"after")')
        assert run.text == "after"
        assert all(run.part is fake_parent.part for run in hyperlink.runs)

    @pytest.mark.parametrize(
        ("value", "exception"),
        [(0, TypeError), (b"label", TypeError), ("bad\x00text", ValueError)],
    )
    def it_rejects_invalid_run_text_before_appending(
        self, value: object, exception: type[Exception], fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element('w:hyperlink/w:r/w:t"before"'))
        hyperlink = Hyperlink(hlink, fake_parent)

        with pytest.raises(exception):
            hyperlink.add_run(cast(str, value))

        assert hlink.xml == xml('w:hyperlink/w:r/w:t"before"')

    def it_does_not_append_a_run_when_its_style_cannot_be_applied(
        self, request: FixtureRequest, fake_parent: t.ProvidesStoryPart
    ):
        property_mock(request, Run, "style", side_effect=KeyError("Missing style"))
        hlink = cast(CT_Hyperlink, element("w:hyperlink"))

        with pytest.raises(KeyError, match="Missing style"):
            Hyperlink(hlink, fake_parent).add_run("label", "Missing style")

        assert hlink.xml == xml("w:hyperlink")

    @pytest.mark.parametrize("value", [None, "", "tip"])
    def it_reads_its_tooltip_from_the_xml_element(
        self, request: FixtureRequest, value: str | None, fake_parent: t.ProvidesStoryPart
    ):
        tooltip_prop = property_mock(request, CT_Hyperlink, "tooltip", return_value=value)
        hyperlink = Hyperlink(cast(CT_Hyperlink, element("w:hyperlink")), fake_parent)

        assert hyperlink.tooltip == value
        tooltip_prop.assert_called_once_with()

    @pytest.mark.parametrize("value", [None, "", "tip"])
    def it_delegates_tooltip_assignment_to_the_xml_element(
        self, request: FixtureRequest, value: str | None, fake_parent: t.ProvidesStoryPart
    ):
        tooltip_prop = property_mock(request, CT_Hyperlink, "tooltip")
        hyperlink = Hyperlink(cast(CT_Hyperlink, element("w:hyperlink")), fake_parent)

        hyperlink.tooltip = value

        tooltip_prop.assert_called_once_with(value)

    @pytest.mark.parametrize(
        ("hlink_cxml", "expected_value"),
        [
            ('w:hyperlink{r:id=rId6}/w:r/w:t"post"', "https://google.com/"),
            ("w:hyperlink{w:anchor=_Toc147925734}", ""),
            ("w:hyperlink", ""),
        ],
    )
    def it_knows_the_hyperlink_address(
        self, hlink_cxml: str, expected_value: str, fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element(hlink_cxml))
        hyperlink = Hyperlink(hlink, fake_parent)

        assert hyperlink.address == expected_value

    @pytest.mark.parametrize(
        ("hlink_cxml", "expected_value"),
        [
            ("w:hyperlink", False),
            ("w:hyperlink/w:r", False),
            ('w:hyperlink/w:r/(w:t"abc",w:lastRenderedPageBreak,w:t"def")', True),
            ('w:hyperlink/w:r/(w:lastRenderedPageBreak,w:t"abc",w:t"def")', True),
            ('w:hyperlink/w:r/(w:t"abc",w:t"def",w:lastRenderedPageBreak)', True),
        ],
    )
    def it_knows_whether_it_contains_a_page_break(
        self, hlink_cxml: str, expected_value: bool, fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element(hlink_cxml))
        hyperlink = Hyperlink(hlink, fake_parent)

        assert hyperlink.contains_page_break is expected_value

    @pytest.mark.parametrize(
        ("hlink_cxml", "expected_value"),
        [
            ("w:hyperlink{r:id=rId6}", ""),
            ("w:hyperlink{w:anchor=intro}", "intro"),
        ],
    )
    def it_knows_the_link_fragment_when_there_is_one(
        self, hlink_cxml: str, expected_value: str, fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element(hlink_cxml))
        hyperlink = Hyperlink(hlink, fake_parent)

        assert hyperlink.fragment == expected_value

    @pytest.mark.parametrize(
        ("hlink_cxml", "count"),
        [
            ("w:hyperlink", 0),
            ("w:hyperlink/w:r", 1),
            ("w:hyperlink/(w:r,w:r)", 2),
            ("w:hyperlink/(w:r,w:lastRenderedPageBreak)", 1),
            ("w:hyperlink/(w:lastRenderedPageBreak,w:r)", 1),
            ("w:hyperlink/(w:r,w:lastRenderedPageBreak,w:r)", 2),
        ],
    )
    def it_provides_access_to_the_runs_it_contains(
        self, hlink_cxml: str, count: int, fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element(hlink_cxml))
        hyperlink = Hyperlink(hlink, fake_parent)

        runs = hyperlink.runs

        actual = [type(item).__name__ for item in runs]
        expected = ["Run" for _ in range(count)]
        assert actual == expected

    @pytest.mark.parametrize(
        ("hlink_cxml", "expected_text"),
        [
            ("w:hyperlink", ""),
            ("w:hyperlink/w:r", ""),
            ('w:hyperlink/w:r/w:t"foobar"', "foobar"),
            ('w:hyperlink/w:r/(w:t"foo",w:lastRenderedPageBreak,w:t"bar")', "foobar"),
            ('w:hyperlink/w:r/(w:t"abc",w:tab,w:t"def",w:noBreakHyphen)', "abc\tdef-"),
        ],
    )
    def it_knows_the_visible_text_of_the_link(
        self, hlink_cxml: str, expected_text: str, fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element(hlink_cxml))
        hyperlink = Hyperlink(hlink, fake_parent)

        text = hyperlink.text

        assert text == expected_text

    @pytest.mark.parametrize(
        ("hlink_cxml", "expected_value"),
        [
            ("w:hyperlink", ""),
            ("w:hyperlink{w:anchor=_Toc147925734}", ""),
            ('w:hyperlink{r:id=rId6}/w:r/w:t"post"', "https://google.com/"),
            (
                'w:hyperlink{r:id=rId6,w:anchor=foo}/w:r/w:t"post"',
                "https://google.com/#foo",
            ),
        ],
    )
    def it_knows_the_full_url_for_web_addresses(
        self, hlink_cxml: str, expected_value: str, fake_parent: t.ProvidesStoryPart
    ):
        hlink = cast(CT_Hyperlink, element(hlink_cxml))
        hyperlink = Hyperlink(hlink, fake_parent)

        assert hyperlink.url == expected_value

    # -- fixtures --------------------------------------------------------------------

    @pytest.fixture
    def fake_parent(self, story_part: Mock, rel: Mock) -> t.ProvidesStoryPart:
        class StoryChild:
            @property
            def part(self) -> StoryPart:
                return story_part

        return StoryChild()

    @pytest.fixture
    def rel(self, request: FixtureRequest):
        return instance_mock(request, _Relationship, target_ref="https://google.com/")

    @pytest.fixture
    def story_part(self, request: FixtureRequest, rel: Mock):
        return instance_mock(request, StoryPart, rels={"rId6": rel})
