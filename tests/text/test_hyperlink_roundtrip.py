"""Saved-document tests for the public hyperlink authoring API."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.styles.style import CharacterStyle
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def paragraph_in(document: DocumentObject, story: str) -> Paragraph:
    if story == "body":
        return document.paragraphs[0]
    if story == "cell":
        return document.tables[0].cell(0, 0).paragraphs[0]
    if story == "header":
        return document.sections[0].header.paragraphs[0]
    if story == "footer":
        return document.sections[0].footer.paragraphs[0]
    raise ValueError(f"Unknown story: {story}")


class DescribeHyperlinkAuthoring:
    @pytest.mark.parametrize("story", ["body", "cell", "header", "footer"])
    @pytest.mark.parametrize(
        "address",
        [
            "https://example.com/a%20b?q=one&lang=en#intro",
            "mailto:hello@example.com?subject=Hello%20there",
            "../guide with spaces.docx",
            "custom:resource",
        ],
    )
    def it_round_trips_formatted_links_in_their_own_story(self, story: str, address: str):
        document = Document()
        document.add_paragraph()
        document.add_table(1, 1)
        paragraph = paragraph_in(document, story)
        paragraph.style = "Heading 1"
        paragraph.add_run("Before ").italic = True
        hyperlink = paragraph.add_hyperlink(address=address, tooltip='Café "tips" & details')
        assert hyperlink.runs == []
        hyperlink.add_run(" Café\t").bold = True
        hyperlink.add_run("code\n", "Emphasis").font.name = "Consolas"
        hyperlink.add_run("tail\r ").italic = True
        paragraph.add_hyperlink("second", address=address)
        paragraph.add_run(" after").bold = True

        stream = BytesIO()
        document.save(stream)
        reopened = Document(stream)
        paragraph = paragraph_in(reopened, story)
        hyperlink, adjacent = paragraph.hyperlinks

        assert paragraph.style is not None
        assert paragraph.style.name == "Heading 1"
        assert paragraph.text == "Before  Café\tcode\ntail\n second after"
        assert [type(item) for item in paragraph.iter_inner_content()] == [
            Run,
            Hyperlink,
            Hyperlink,
            Run,
        ]
        assert [run.text for run in paragraph.runs] == ["Before ", " after"]
        assert paragraph.runs[0].italic is True
        assert paragraph.runs[1].bold is True
        assert hyperlink.url == hyperlink.address == adjacent.url == address
        assert hyperlink.fragment == ""
        assert hyperlink.tooltip == 'Café "tips" & details'
        assert adjacent.tooltip is None
        assert hyperlink.runs[0].bold is True
        assert hyperlink.runs[1].style.name == "Emphasis"
        assert hyperlink.runs[1].font.name == "Consolas"
        assert hyperlink.runs[2].italic is True
        assert all(run.part is paragraph.part for run in hyperlink.runs)

        ids = paragraph.part.element.xpath(".//w:hyperlink/@r:id")
        assert ids[0] == ids[1]
        rel = paragraph.part.rels[ids[0]]
        assert rel.reltype == RT.HYPERLINK
        assert rel.is_external
        assert rel.target_ref == address
        if story in ("header", "footer"):
            assert not any(rel.reltype == RT.HYPERLINK for rel in reopened.part.rels.values())

    @pytest.mark.parametrize("story", ["body", "cell", "header", "footer"])
    def it_can_add_a_picture_to_a_hyperlink_run(self, story: str):
        document = Document()
        document.add_paragraph()
        document.add_table(1, 1)
        paragraph = paragraph_in(document, story)
        hyperlink = paragraph.add_hyperlink(address="https://example.com")
        image_path = Path(__file__).parents[1] / "test_files" / "python-icon.jpeg"
        picture = hyperlink.add_run().add_picture(str(image_path))
        assert picture.width > 0

        stream = BytesIO()
        document.save(stream)
        paragraph = paragraph_in(Document(stream), story)
        assert paragraph.hyperlinks[0].url == "https://example.com"
        embeds = paragraph.part.element.xpath(".//w:hyperlink/w:r/w:drawing//a:blip/@r:embed")
        assert len(embeds) == 1
        rel = paragraph.part.rels[embeds[0]]
        assert rel.reltype == RT.IMAGE
        assert not rel.is_external
        assert rel.target_part.blob == image_path.read_bytes()

    def it_leaves_style_definitions_under_caller_control(self):
        document = Document()
        assert "Hyperlink" not in document.styles
        link = document.add_paragraph().add_hyperlink("plain", address="guide.pdf")
        assert "Hyperlink" not in document.styles
        assert link.runs[0].style.name == "Default Paragraph Font"
        style = document.styles.add_style(  # pyright: ignore[reportUnknownMemberType]
            "Hyperlink", WD_STYLE_TYPE.CHARACTER
        )
        assert isinstance(style, CharacterStyle)
        style.font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_2
        style.font.underline = False
        link.add_run("styled", style)
        link.add_run("also styled", "Hyperlink")
        link.add_run("override", "Emphasis")

        stream = BytesIO()
        document.save(stream)
        reopened = Document(stream)

        style = reopened.styles["Hyperlink"]
        assert isinstance(style, CharacterStyle)
        assert style.font.color.theme_color == MSO_THEME_COLOR_INDEX.ACCENT_2
        assert style.font.underline is False
        assert [run.style.name for run in reopened.paragraphs[0].hyperlinks[0].runs] == [
            "Default Paragraph Font",
            "Hyperlink",
            "Hyperlink",
            "Emphasis",
        ]

    @pytest.mark.parametrize("value", [None, "", 'Café "tips" & details'])
    def it_round_trips_tooltip_replacement(self, value: str | None):
        document = Document()
        hyperlink = document.add_paragraph().add_hyperlink(
            "label", address="guide.pdf", tooltip="old"
        )
        hyperlink.tooltip = value

        stream = BytesIO()
        document.save(stream)

        assert Document(stream).paragraphs[0].hyperlinks[0].tooltip == value
