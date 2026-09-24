"""Internal hyperlink authoring through the public paragraph API."""

from io import BytesIO

import pytest

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT


class DescribeInternalHyperlinks:
    def it_supports_forward_references_without_external_relationships(self):
        document = Document()
        paragraph = document.add_paragraph()
        link = paragraph.add_hyperlink("Jump ", anchor="Target", tooltip="Details")
        link.add_run("there").bold = True
        document.bookmarks.add("Target", paragraph=document.add_paragraph("Target heading"))
        stream = BytesIO()
        document.save(stream)
        loaded = Document(stream)
        link = loaded.paragraphs[0].hyperlinks[0]
        assert link.address == ""
        assert link.fragment == "Target"
        assert link.text == "Jump there"
        assert link.tooltip == "Details"
        assert link.runs[1].bold is True
        assert not any(rel.reltype == RT.HYPERLINK for rel in loaded.part.rels.values())

    @pytest.mark.parametrize(
        "kwargs",
        [
            {"anchor": ""},
            {"anchor": "a b"},
            {"anchor": "#Target"},
            {"anchor": "Target", "address": "https://example.com"},
            {"anchor": "Target", "text": "bad\x00text"},
        ],
    )
    def it_validates_before_mutation(self, kwargs):
        document = Document()
        paragraph = document.add_paragraph("Before")
        before = paragraph._p.xml
        relationships = list(document.part.rels)
        with pytest.raises(ValueError, match="bookmark|exactly one|XML"):
            paragraph.add_hyperlink(**kwargs)
        assert paragraph._p.xml == before
        assert list(document.part.rels) == relationships
